"""Build presentation.html (slides as images plus speaker notes) from the exported PDF and the .pptx.

Usage: python make_viewer.py <deck.pdf> <deck.pptx> <out_dir> <viewer_template.html>
Writes <out_dir>/presentation.html and <out_dir>/slides/NN.jpg.
"""
import sys, json, re, html
from pathlib import Path
import fitz
from pptx import Presentation

pdf, pptx, out = sys.argv[1], sys.argv[2], Path(sys.argv[3])
(out / "slides").mkdir(parents=True, exist_ok=True)
doc = fitz.open(pdf)
for i, page in enumerate(doc, 1):
    zoom = 1920 / page.rect.width
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
    pix.save(str(out / "slides" / f"{i:02d}.jpg"), jpg_quality=90)
p = Presentation(pptx)
sections = [(1, 2, "Opening"), (3, 7, "The problem"), (8, 9, "The method"), (10, 15, "The findings"), (16, 20, "What it means"), (21, 25, "Backup")]
def section(i):
    for a, b, n in sections:
        if a <= i <= b:
            return n
slides = []
for i, s in enumerate(p.slides, 1):
    title = s.shapes.title.text_frame.text.strip() if s.shapes.title is not None else ""
    note = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
    m = re.match(r"^\[([^\]]*)\]\s*(.*)$", note, re.S)
    timing, body = (m.group(1), m.group(2)) if m else ("", note)
    # a parenthetical after the stamp (e.g. "(shorten ...)") stays with the timing
    m2 = re.match(r"^\(([^)]*)\)\s*(.*)$", body, re.S)
    hint = ""
    if m2:
        hint, body = m2.group(1), m2.group(2)
    lead = re.match(r"^(Leave this slide up[^.\n]*\.)\s*(.*)$", body, re.S)
    if lead:
        hint, body = lead.group(1), lead.group(2)
    paras = [x.strip() for x in body.split("\n") if x.strip()]
    slides.append({"n": i, "title": title.replace("\n", " "), "section": section(i), "timing": timing, "hint": hint, "notes": paras})
data = json.dumps(slides, ensure_ascii=False).replace("</", "<\\/")
main = [" ".join(q for q in x["notes"] if not q.startswith("(Backup")) for x in slides if x["n"] <= 20]
words = sum(len(re.sub(r"\[pause[^\]]*\]", "", t).split()) for t in main)
pauses = sum(int(m) if m else 2 for t in main for m in re.findall(r"\[pause(?:, (\d+) s)?\]", t))
mins = (words / 135 * 60 + pauses + 3 * len(main)) / 60
tpl = Path(sys.argv[4]).read_text()
page = tpl.replace("__DATA__", data).replace("__WORDS__", str(words)).replace("__MINUTES__", f"{mins:.1f}")
(out / "presentation.html").write_text(page)
print("slides:", len(slides), "| main-talk words:", words, "| html bytes:", len(page.encode()))
