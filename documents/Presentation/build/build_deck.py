"""Build the thesis presentation on the official KTH PowerPoint template (Dec 2025).

Run with a Python that has python-pptx:  python build_deck.py
Reads deck_data.json (produced by deck_data.py) so every number traces back to the
thesis analysis outputs. Writes ../Dante_Wesslund_Thesis_Presentation.pptx.
"""
import copy
import json
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
D = json.load(open(HERE / "deck_data.json"))
OUT = HERE.parent / "Dante_Wesslund_Thesis_Presentation.pptx"

# KTH palette (template custom colours)
NAVY, KTH, SKY, LIGHT, SAND = "000061", "004791", "6298D2", "DEF0FF", "EBE5E0"
GREY, LGREY, DGREY = "A5A5A5", "E6E6E6", "323232"
GREEN, LGREEN, BRICK, LBRICK, LYELLOW, DYELLOW = "4DA060", "C7EBBA", "E86A58", "FFCCC4", "FFF0B0", "A65900"
WHITE = "FFFFFF"
NAIVE_C, MIX_C = SKY, NAVY          # retrieval modes, used consistently
TFIDF_C, MB_C, STATIC_C = SKY, KTH, GREY

X0, W_FULL = 0.66, 12.02             # content frame (inches)
CONTENT_TOP = 2.2

def rgb(h):
    return RGBColor.from_string(h)

def fmt_int(x):
    return f"{int(x):,}"

def pv(p):
    return "p < 0.001" if p < 0.001 else (f"p = {p:.3f}" if p < 0.2 else f"p = {p:.2f}")

# ----------------------------------------------------------------------------------------------
prs = Presentation(HERE / "kth_template_251203.pptx")
LAYOUT = {l.name.strip(): l for l in prs.slide_masters[0].slide_layouts}
sld = prs.slides._sldIdLst
for s in list(sld):
    prs.part.drop_rel(s.rId)
    sld.remove(s)

FOOTER = "Adaptive Hybrid RAG  ·  Dante Wesslund  ·  KTH 2026"

def add_footer(slide, footer=True):
    for ph in slide.slide_layout.placeholders:
        t = ph.placeholder_format.type
        if t == PP_PLACEHOLDER.SLIDE_NUMBER or (footer and t == PP_PLACEHOLDER.FOOTER):
            slide.shapes._spTree.append(copy.deepcopy(ph._element))
    for ph in slide.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.FOOTER:
            ph.text_frame.text = FOOTER

def alt(shape, text):
    shape._element.xpath(".//p:cNvPr")[0].set("descr", text)

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()

def set_runs(tf, paras, size=22, color=DGREY, bold=False, align=PP_ALIGN.LEFT, space_after=6, line=1.05, bullets=False):
    """paras: list of str | list of (text, {opts}) runs | dict(runs=..., size=..., ...)."""
    tf.word_wrap = True
    first = True
    for para in paras:
        popt = {}
        if isinstance(para, dict):
            popt = para
            runs = para["runs"]
        else:
            runs = para
        if isinstance(runs, str):
            runs = [(runs, {})]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = popt.get("align", align)
        p.space_after = Pt(popt.get("space_after", space_after))
        p.line_spacing = popt.get("line", line)
        if popt.get("bullet", bullets):
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.28))))
            pPr.set("indent", str(-Emu(Inches(0.28))))
            for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
                for e in pPr.findall(qn(tag)):
                    pPr.remove(e)
            bf = etree.SubElement(pPr, qn("a:buFont")); bf.set("typeface", "Arial")
            bc = etree.SubElement(pPr, qn("a:buChar")); bc.set("char", "•")
        for text, o in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.size = Pt(o.get("size", popt.get("size", size)))
            f.bold = o.get("bold", popt.get("bold", bold))
            f.italic = o.get("italic", False)
            f.color.rgb = rgb(o.get("color", popt.get("color", color)))
    return tf

def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0.0, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    set_runs(tf, paras, **kw)
    return tb

def box(slide, x, y, w, h, fill, paras=None, radius=0.06, line=None, anchor=MSO_ANCHOR.TOP, pad=0.18, shape=MSO_SHAPE.ROUNDED_RECTANGLE, **kw):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(pad))
    if paras:
        set_runs(tf, paras, **kw)
    return s

def arrow(slide, x1, y1, x2, y2, color=GREY, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c

def line(slide, x1, y1, x2, y2, color=GREY, width=1.5, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(width)
    if dash:
        ln = c.line._get_or_add_ln()
        pd = etree.SubElement(ln, qn("a:prstDash")); pd.set("val", dash)
    return c

def dot(slide, cx, cy, d, fill, outline=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(WHITE)
    if outline:
        s.line.color.rgb = rgb(outline); s.line.width = Pt(2)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s

def content_slide(title, section_note=None):
    s = prs.slides.add_slide(LAYOUT["Headline only"])
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(X0), Inches(0.95), Inches(W_FULL), Inches(1.05)
    tf = t.text_frame
    tf.text = title
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.word_wrap = True
    for p in tf.paragraphs:
        p.line_spacing = 0.95
        for r in p.runs:
            r.font.size = Pt(30)
    add_footer(s)
    return s

def chapter_slide(title, sub, layout="Title/chapter, KTH blue"):
    s = prs.slides.add_slide(LAYOUT[layout])
    s.shapes.title.text_frame.text = title
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 1:
            ph.text_frame.text = sub
    add_footer(s, footer=False)
    return s

def chart_base(chart, size=16):
    chart.font.size = Pt(size)
    chart.font.color.rgb = rgb(DGREY)
    cs = chart._chartSpace
    sp = cs.find(qn("c:spPr"))
    if sp is None:
        sp = etree.SubElement(cs, qn("c:spPr"))
        txPr = cs.find(qn("c:txPr"))
        if txPr is not None:
            cs.remove(sp); txPr.addprevious(sp)
    for ch in list(sp):
        sp.remove(ch)
    etree.SubElement(sp, qn("a:noFill"))
    ln = etree.SubElement(sp, qn("a:ln")); etree.SubElement(ln, qn("a:noFill"))

def quiet_axes(chart, vmax=None, vmin=0, fmt="0.0", grid=True, show_val=True, size=15):
    va = chart.value_axis
    if vmax is not None:
        va.maximum_scale = vmax
    va.minimum_scale = vmin
    va.has_major_gridlines = grid
    if grid:
        va.major_gridlines.format.line.color.rgb = rgb(LGREY)
        va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(size)
    va.tick_labels.number_format = fmt
    va.tick_labels.number_format_is_linked = False
    if not show_val:
        va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    ca = chart.category_axis
    ca.format.line.color.rgb = rgb(GREY)
    ca.tick_labels.font.size = Pt(size + 1)
    ca.has_major_gridlines = False

def color_series(series, hexcol):
    series.format.fill.solid(); series.format.fill.fore_color.rgb = rgb(hexcol)
    series.format.line.fill.background()

def bar_labels(plot, fmt="0.000", size=14, pos=XL_LABEL_POSITION.OUTSIDE_END, color=DGREY):
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = fmt; dl.number_format_is_linked = False
    dl.position = pos
    dl.font.size = Pt(size); dl.font.color.rgb = rgb(color)

def img(slide, name, x, y, w, h, descr):
    pic = slide.shapes.add_picture(str(HERE / "charts" / f"{name}.png"), Inches(x), Inches(y), Inches(w), Inches(h))
    alt(pic, descr)
    return pic

P = D["policies"]; CMP = D["cmp"]; MC = D["mcnemar"]; TM = D["token_means"]
mb, tf_, ty, an, am = P["modernbert"], P["tfidf"], P["type_only"], P["always_naive"], P["always_mix"]


NB = " "   # non-breaking space, keeps "p = 0.003" on one line
MINUS = "−"

def sgn(x, nd=3):
    """Signed number with a true minus sign."""
    return f"{x:+.{nd}f}".replace("-", MINUS)

def num(x, nd=3):
    return f"{x:.{nd}f}".replace("-", MINUS)

def pnb(p):
    """p-value phrase with non-breaking spaces."""
    if p < 0.001:
        return f"p{NB}<{NB}0.001"
    return f"p{NB}={NB}" + (f"{p:.3f}" if p < 0.2 else f"{p:.2f}")

TOK = {k: P[k]["tokens_printed"] for k in P}
CAS = D["cascade"]
bt = D["by_type"]
n_ans = D["classes"]["naive_enough"] + D["classes"]["mix_required"]

# ============================================================================================ 1
s = prs.slides.add_slide(LAYOUT["Title, KTH blue"])
t = s.shapes.title
t.text_frame.text = "Adaptive Hybrid Retrieval-Augmented Generation"
t.left, t.top, t.width, t.height = Inches(1.85), Inches(3.15), Inches(9.62), Inches(1.85)
for ph in s.placeholders:
    if ph.placeholder_format.idx == 1:
        ph.text_frame.text = "Dante Wesslund  ·  KTH EECS  ·  22 October 2026"
        ph.left, ph.top, ph.width, ph.height = Inches(2.33), Inches(6.12), Inches(8.67), Inches(0.51)
text(s, 2.4, 5.2, 8.53, 0.78, [{"runs": "Optimizing the computational cost-performance trade-off", "align": PP_ALIGN.CENTER, "size": 20, "color": WHITE, "space_after": 0},
                                {"runs": "via learned query routing", "align": PP_ALIGN.CENTER, "size": 20, "color": WHITE}])
text(s, 2.9, 6.72, 7.53, 0.6, [{"runs": "Supervisors: Arvid Eriksson (KTH), Peng Zhang (Ericsson)", "align": PP_ALIGN.CENTER, "size": 14, "color": "DEF0FF", "space_after": 0},
                                {"runs": "Examiner: Pawel Herman  ·  In collaboration with Ericsson", "align": PP_ALIGN.CENTER, "size": 14, "color": "DEF0FF"}])
notes(s, """
[0:00 to 0:30]
Hello everyone. Before I start, can everyone see the title slide?
My name is Dante Wesslund, and this is my master's thesis, Adaptive Hybrid Retrieval-Augmented Generation, done at Ericsson. Thank you to my supervisors, Arvid Eriksson at KTH and Peng Zhang at Ericsson, and to my examiner Pawel Herman.
In one sentence: can a small, cheap classifier look at a question and decide, before any retrieval runs, whether the cheap or the expensive retrieval mode is needed, and does that save real cost?
""")

# ============================================================================================ 2
s = content_slide("Graph-enhanced retrieval answers harder questions, but every query pays for it")
cw, gap = 5.45, 1.12
xL, xR = X0, X0 + cw + gap
yc, hc = 2.3, 3.3
box(s, xL, yc, cw, hc, LIGHT, [
    {"runs": [("Naive", {"bold": True, "size": 26, "color": NAVY}), ("   vector search only", {"size": 20, "color": DGREY})], "space_after": 14},
    {"runs": [(f"{an['mean_time']:.1f} s", {"bold": True, "size": 36, "color": KTH})], "space_after": 0},
    {"runs": "mean time per query", "size": 18, "space_after": 12},
    {"runs": [(f"≈{fmt_int(round(TM['naive'], -2))}", {"bold": True, "size": 36, "color": KTH})], "space_after": 0},
    {"runs": "model tokens per query*", "size": 18},
], pad=0.3)
box(s, xR, yc, cw, hc, NAVY, [
    {"runs": [("Mix", {"bold": True, "size": 26, "color": WHITE}), ("   vector + knowledge graph", {"size": 20, "color": LIGHT})], "space_after": 14},
    {"runs": [(f"{am['mean_time']:.1f} s", {"bold": True, "size": 36, "color": WHITE})], "space_after": 0},
    {"runs": "mean time per query", "size": 18, "color": LIGHT, "space_after": 12},
    {"runs": [(f"≈{fmt_int(round(TM['mix'], -2))}", {"bold": True, "size": 36, "color": WHITE})], "space_after": 0},
    {"runs": "model tokens per query*", "size": 18, "color": LIGHT},
], pad=0.3)
ratio_t = am["mean_time"] / an["mean_time"]; ratio_k = TM["mix"] / TM["naive"]
text(s, xL + cw, yc + 0.85, gap, 0.9, [{"runs": f"{ratio_t:.1f}×", "align": PP_ALIGN.CENTER, "size": 22, "bold": True, "color": NAVY, "space_after": 0},
                                       {"runs": "time", "align": PP_ALIGN.CENTER, "size": 15, "color": "595959"}])
text(s, xL + cw, yc + 2.05, gap, 0.9, [{"runs": f"{ratio_k:.1f}×", "align": PP_ALIGN.CENTER, "size": 22, "bold": True, "color": NAVY, "space_after": 0},
                                       {"runs": "tokens", "align": PP_ALIGN.CENTER, "size": 15, "color": "595959"}])
mr, ne, nj = D["classes"]["mix_required"], D["classes"]["naive_enough"], D["n_judged"]
box(s, X0, 5.78, W_FULL, 0.78, SAND, [{"runs": [
    ("Yet Mix was needed on only ", {}), (f"{mr} of {fmt_int(nj)}", {"bold": True, "color": NAVY}),
    (f" judged questions ({mr/nj*100:.0f}%). On {ne} ({ne/nj*100:.0f}%), Naive was already enough.", {})], "size": 20}],
    anchor=MSO_ANCHOR.MIDDLE, pad=0.25)
text(s, X0, 6.62, W_FULL, 0.3, [{"runs": f"Times: all {D['n_test']} test questions.  *Tokens: per-mode means over {TM['n']} instrumented queries (prompt, output and reasoning tokens).",
                                "size": 13, "color": "595959"}])
notes(s, f"""
[0:30 to 1:30]
Retrieval-augmented generation grounds a language model in documents. The cheapest form is vector search: embed the question and fetch the most similar chunks. In LightRAG that is the Naive mode.
Graph-enhanced retrieval adds an entity-relation graph, so the system can connect facts that live in different documents. That is LightRAG's Mix mode, and it helps on multi-hop questions.
But it is expensive. In my runs Mix takes {am['mean_time']:.1f} seconds per query against {an['mean_time']:.1f} for Naive, about {ratio_t:.1f} times longer, and it uses roughly {ratio_k:.0f} times as many model tokens.
And here is the key observation: of the {fmt_int(nj)} questions I judged, Mix was needed on only {mr}, about one in six. On almost half, {ne}, Naive was already enough. So a system that always runs Mix pays the graph premium on every question, but it changes the outcome on only about one in six.
(The token numbers come from a 15-query instrumented subset. I return to that at the end.)
""")

# ============================================================================================ 3
s = content_slide("Some questions need the graph, most do not")
cw = 5.86
def example_card(x, q, gold, naive_ok, naive_txt, mix_txt, verdict, verdict_col):
    box(s, x, 2.3, cw, 4.25, WHITE, line=LGREY, radius=0.04)
    text(s, x + 0.25, 2.45, cw - 0.5, 0.95, [{"runs": [("Q  ", {"bold": True, "color": KTH}), (q, {"bold": True, "color": NAVY})], "size": 20, "line": 1.0}])
    text(s, x + 0.25, 3.38, cw - 0.5, 0.35, [{"runs": [("Gold answer: ", {"color": "595959"}), (gold, {"bold": True})], "size": 17}])
    for yy, mode, ok, t in ((3.85, "Naive", naive_ok, naive_txt), (5.0, "Mix", True, mix_txt)):
        box(s, x + 0.25, yy, 1.3, 0.4, GREEN if ok else BRICK, [{"runs": "correct" if ok else "not enough", "align": PP_ALIGN.CENTER, "size": 14, "bold": True, "color": WHITE}],
            radius=0.3, pad=0.02, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 1.7, yy - 0.03, cw - 2.0, 1.05, [{"runs": [(mode + ":  ", {"bold": True, "color": NAIVE_C if mode == "Naive" else MIX_C}), (t, {"italic": True})], "size": 15, "line": 1.0}])
    box(s, x + 0.25, 5.95, cw - 0.5, 0.48, verdict_col, [{"runs": verdict, "size": 17, "bold": True, "color": WHITE if verdict_col == NAVY else NAVY}],
        anchor=MSO_ANCHOR.MIDDLE, pad=0.15, radius=0.2)
example_card(X0, "Which film came out first, The Love Route or Engal Aasan?", "The Love Route", True,
             "“…The Love Route came out significantly earlier than Engal Aasan.” (1915 vs 2009)",
             "“…The Love Route was released significantly earlier than Engal Aasan.”",
             "Label: naive_enough  →  route to Naive", LIGHT)
example_card(X0 + cw + 0.3, "When was the director of film Private Hell 36 born?", "October 26, 1912", False,
             "“…was directed by Don Siegel. However, the provided documents do not contain information regarding Don Siegel's birth date.”",
             "“…Donald Siegel was born on October 26, 1912.”",
             "Label: mix_required  →  route to Mix", NAVY)
text(s, X0, 6.62, W_FULL, 0.3, [{"runs": "Real questions and answers from the 2WikiMultihopQA runs (abridged). Comparison question left, two-hop compositional question right.", "size": 13, "color": "595959"}])
notes(s, """
[1:30 to 2:40]
Two real examples from my runs make the difference concrete.
On the left: which film came out first, The Love Route or Engal Aasan? Each film has its own document with a release date, vector search fetches both, and both modes answer correctly. Mix adds nothing here, so the label is naive_enough and the right decision is the cheap mode.
On the right: when was the director of Private Hell 36 born? This is a two-hop question. Vector search finds the film's page and even names the director, Don Siegel, but his birth date is in a different document that does not look similar to the question. Naive gives up. Mix's graph context links the film to the director entity, and it answers correctly: October 26, 1912. Here the expensive mode is needed.
The whole thesis is about telling these two situations apart from the question alone.
""")

# ============================================================================================ 4
s = content_slide("The idea: let a cheap classifier pick the mode from the question, before retrieval runs")
fy = 3.55
box(s, X0, fy, 1.9, 1.0, LIGHT, [{"runs": "Question", "align": PP_ALIGN.CENTER, "bold": True, "size": 20, "color": NAVY}], anchor=MSO_ANCHOR.MIDDLE)
box(s, X0 + 2.45, fy - 0.35, 2.2, 1.7, KTH, [{"runs": "Router", "align": PP_ALIGN.CENTER, "bold": True, "size": 22, "color": WHITE, "space_after": 2},
                                             {"runs": "small classifier,\nno LLM call", "align": PP_ALIGN.CENTER, "size": 15, "color": LIGHT}], anchor=MSO_ANCHOR.MIDDLE)
arrow(s, X0 + 1.9, fy + 0.5, X0 + 2.45, fy + 0.5)
box(s, X0 + 5.2, 2.55, 1.95, 0.95, NAIVE_C, [{"runs": "Naive", "align": PP_ALIGN.CENTER, "bold": True, "size": 20, "color": WHITE, "space_after": 0},
                                              {"runs": "low cost", "align": PP_ALIGN.CENTER, "size": 14, "color": WHITE}], anchor=MSO_ANCHOR.MIDDLE)
box(s, X0 + 5.2, 4.6, 1.95, 0.95, MIX_C, [{"runs": "Mix", "align": PP_ALIGN.CENTER, "bold": True, "size": 20, "color": WHITE, "space_after": 0},
                                            {"runs": "high cost", "align": PP_ALIGN.CENTER, "size": 14, "color": LIGHT}], anchor=MSO_ANCHOR.MIDDLE)
arrow(s, X0 + 4.65, fy + 0.25, X0 + 5.2, 3.02)
arrow(s, X0 + 4.65, fy + 0.75, X0 + 5.2, 5.07)
text(s, X0, 5.95, 7.2, 0.8, [{"runs": "Trained on labels that say which mode was enough for each question (next section).", "size": 17, "color": "595959"}])
rx, rw = 8.55, 4.13
for i, (big, small) in enumerate((
        (f"{MINUS}{mb['time_saving']*100:.0f}%", "mean time per query vs. always running Mix"),
        (f"{MINUS}{mb['token_saving']*100:.0f}%", "estimated model tokens per query"),
        (f"≈{mb['retention']*100:.0f}%", "of always-Mix's judged-correct answers kept*"))):
    yy = 2.3 + i * 1.42
    box(s, rx, yy, rw, 1.25, LIGHT if i < 2 else SAND, [
        {"runs": [(big, {"bold": True, "size": 38, "color": KTH})], "space_after": 0},
        {"runs": small, "size": 16, "line": 1.0}], anchor=MSO_ANCHOR.MIDDLE, pad=0.22)
text(s, rx, 6.58, rw, 0.4, [{"runs": "ModernBERT router on 193 test questions. *Assumes Mix is correct wherever Naive was.", "size": 12, "color": "595959"}])
notes(s, f"""
[2:40 to 3:40]
So the idea is to put a router in front of retrieval. It reads only the question, no retrieved text and no language-model call, and sends it to Naive or Mix.
The router is trained on labels that record, for every question, which mode was actually enough. I will show how those labels are made in a moment.
Here is where we end up, so you know what to listen for. My best router, a fine-tuned ModernBERT, cuts mean time per query by {mb['time_saving']*100:.0f} percent compared with always running Mix, cuts estimated model tokens by {mb['token_saving']*100:.0f} percent, and keeps about {mb['retention']*100:.0f} percent of the answers that always-Mix gets right, under an assumption I will make explicit.
The rest of the talk is about how solid those numbers are.
""")

# ============================================================================================ 5
s = content_slide("Can a lightweight classifier pick the cheapest sufficient retrieval mode?")
text(s, X0, 2.25, W_FULL, 1.0, [{"runs": "Within one fixed Hybrid RAG setting (LightRAG, Naive vs. Mix), how well can a query-only classifier predict the minimum sufficient retrieval mode, and what does routing with it save?",
                                  "size": 21, "color": DGREY, "line": 1.05}])
cw = 5.86
for i, (n_, head, body) in enumerate((
        ("1", "Prediction", "How does a fine-tuned ModernBERT classifier compare to a TF-IDF baseline and a question-type heuristic in predicting the label?"),
        ("2", "Cost", "How much mean time and model-token cost does routing save compared with always-Mix, and what does it risk in answer quality?"))):
    x = X0 + i * (cw + 0.3)
    box(s, x, 3.65, cw, 2.75, LIGHT if i == 0 else SAND, [
        {"runs": [(f"Sub-question {n_}  ", {"bold": True, "color": KTH, "size": 18})], "space_after": 4},
        {"runs": head, "bold": True, "size": 26, "color": NAVY, "space_after": 10},
        {"runs": body, "size": 20, "line": 1.05}], pad=0.3)
notes(s, """
[3:40 to 4:25]
Formally, the research question is: within one fixed Hybrid RAG setting, how well can a lightweight query-only classifier predict the minimum sufficient retrieval mode, and what does routing with it save?
Why only the question? Because the point is to decide before any retrieval cost is paid. Anything that inspects retrieved text has already paid for at least one retrieval run.
I split it into two sub-questions. First, prediction: how does a fine-tuned encoder, ModernBERT, compare to a simple TF-IDF baseline, and to a heuristic that uses the benchmark's question type? Second, cost: how much time and model-token cost does routing save relative to always running Mix, and what does it risk in answer quality?
""")

# ============================================================================================ 6
s = content_slide("Prior work routes between models or retrieval depth; this thesis routes between vector and graph retrieval")
rows = (
    ("Adaptive-RAG", "Jeong et al., 2024", "A question-complexity classifier picks no, single-step or multi-step retrieval."),
    ("RouterBench", "Hu et al., 2024", "Routes each query between language models on cost vs. quality."),
    ("GraphRAG, LightRAG", "Edge et al., 2025; Guo et al., 2025", "Graph-enhanced retrieval improves multi-hop answers at higher cost."),
)
yy = 2.3
for name, cite, what in rows:
    text(s, X0, yy, 3.3, 0.8, [{"runs": name, "bold": True, "size": 20, "color": NAVY, "space_after": 0}, {"runs": cite, "size": 14, "color": "595959"}])
    text(s, X0 + 3.5, yy + 0.05, W_FULL - 3.5, 0.8, [{"runs": what, "size": 20}])
    line(s, X0, yy + 0.95, X0 + W_FULL, yy + 0.95, LGREY, 1)
    yy += 1.08
box(s, X0, yy + 0.1, W_FULL, 1.15, KTH, [{"runs": [("This thesis  ", {"bold": True, "size": 20, "color": WHITE}),
    ("A query-only router choosing between two retrieval modes inside one LightRAG index, trained on judge-derived labels and evaluated with paired statistical tests.", {"size": 20, "color": WHITE})]}],
    anchor=MSO_ANCHOR.MIDDLE, pad=0.3)
notes(s, """
[4:25 to 5:10]  (shorten to one sentence if running long)
Where does this sit? Adaptive-RAG is the closest precedent: a small classifier estimates question complexity and picks how much retrieval to do. RouterBench frames routing between language models as a cost-quality trade-off. GraphRAG and LightRAG show that graph retrieval helps on multi-hop questions, and also document its cost.
What none of them asks is how to choose between a vector mode and a graph-enhanced mode inside one framework, from the question alone. That is the gap I study, with labels derived by an LLM judge and paired statistical tests on the results.
""")

# ============================================================================================ 7
s = chapter_slide("Method", "From benchmark questions to labels and routers")
notes(s, "[5:10]  Moving on to the method: how I turned benchmark questions into training labels, which routers I trained, and how I evaluated them.")

# ============================================================================================ 8
s = content_slide("Both modes run on every question, and an LLM judge records which one was enough")
bw, bh = 2.62, 1.45
stepx = (W_FULL - 4 * bw) / 3
row1 = (("2,000 questions", "2WikiMultihopQA, multi-hop by design", LIGHT, NAVY),
        ("LightRAG index", "shared corpus: vector chunks + entity-relation graph", LIGHT, NAVY),
        ("Run both modes", "Naive and Mix answers, time per query", LIGHT, NAVY),
        ("LLM-as-a-Judge", "compares both answers with the gold answer", KTH, WHITE))
text(s, X0, 2.2, 6, 0.35, [{"runs": "LABELLING  (offline, once)", "size": 14, "bold": True, "color": KTH}])
for i, (h, b, fill, col) in enumerate(row1):
    x = X0 + i * (bw + stepx)
    box(s, x, 2.6, bw, bh, fill, [{"runs": h, "bold": True, "size": 20, "color": col, "space_after": 4},
                                  {"runs": b, "size": 15, "color": col if col == WHITE else DGREY, "line": 1.0}], anchor=MSO_ANCHOR.MIDDLE, pad=0.18)
    if i < 3:
        arrow(s, x + bw + 0.05, 2.6 + bh / 2, x + bw + stepx - 0.05, 2.6 + bh / 2)
text(s, X0, 4.35, 6, 0.35, [{"runs": "LEARNING AND EVALUATION", "size": 14, "bold": True, "color": KTH}])
row2 = (("Binary dataset", f"{fmt_int(D['n_binary'])} questions: {D['split_n']['train']} train / {D['split_n']['validation']} val / {D['split_n']['test']} test", SAND, NAVY),
        ("Train routers", "TF-IDF + logistic regression; ModernBERT (fine-tuned)", SAND, NAVY),
        ("Evaluate", "as classifiers and as routing policies on cost", SAND, NAVY))
bw2 = (W_FULL - 2 * stepx) / 3
for i, (h, b, fill, col) in enumerate(row2):
    x = X0 + i * (bw2 + stepx)
    box(s, x, 4.75, bw2, 1.3, fill, [{"runs": h, "bold": True, "size": 20, "color": col, "space_after": 4},
                                     {"runs": b, "size": 15, "line": 1.0}], anchor=MSO_ANCHOR.MIDDLE, pad=0.18)
    if i < 2:
        arrow(s, x + bw2 + 0.05, 5.4, x + bw2 + stepx - 0.05, 5.4)
text(s, X0, 6.3, W_FULL, 0.5, [{"runs": "Same Gemini model generates and judges, temperature 0, 5 text chunks per mode; first 2,000 queryable questions of the 2WikiMultihopQA train split, in benchmark order.",
                                 "size": 14, "color": "595959"}])
notes(s, """
[5:10 to 6:25]
This is the whole pipeline. I took 2,000 questions from 2WikiMultihopQA, a public benchmark whose questions need several documents by construction. Their documents are indexed once in LightRAG, which builds both the vector index and the entity-relation graph.
Every question is then run through both Naive and Mix, so for each question I have two answers and two execution times.
Then an LLM judge, Gemini, gets the question, the gold answer from the benchmark, and both candidate answers, and decides which mode was enough. Giving it the gold answer turns judging into a constrained comparison rather than open-ended scoring.
This labelling is done once, offline. The bottom row is the learning part: build a binary dataset, train the routers, and evaluate them both as classifiers and as routing policies.
""")

# ============================================================================================ 9
s = content_slide("The label is the cheapest mode that was judged sufficient")
tot = D["n_judged"]
segs = ((D["classes"]["naive_enough"], NAIVE_C, WHITE, "Naive enough"),
        (D["classes"]["mix_required"], MIX_C, WHITE, "Mix"),
        (D["classes"]["none_enough"], LGREY, DGREY, "Neither correct"))
xx = X0
for n_, fill, col, lab in segs:
    w = W_FULL * n_ / tot
    box(s, xx, 2.35, w - 0.04, 1.05, fill, [{"runs": [(f"{fmt_int(n_)}  ", {"bold": True, "size": 24, "color": col}), (lab, {"size": 18, "color": col})]}],
        anchor=MSO_ANCHOR.MIDDLE, pad=0.15, radius=0.04, shape=MSO_SHAPE.RECTANGLE)
    xx += w
text(s, X0, 3.45, W_FULL, 0.35, [{"runs": f"{fmt_int(tot)} judged questions", "size": 15, "color": "595959"}])
rules = (("Naive answer correct", "naive_enough", "label 0", NAIVE_C),
         ("Naive not enough, Mix correct", "mix_required", "label 1", MIX_C),
         ("Neither correct: routing cannot help", "none_enough", "excluded", GREY))
yy = 4.05
for cond, cls, lab, col in rules:
    dot(s, X0 + 0.15, yy + 0.22, 0.24, col)
    text(s, X0 + 0.45, yy, 4.35, 0.45, [{"runs": cond, "size": 19}])
    text(s, X0 + 4.9, yy, 2.1, 0.45, [{"runs": cls, "size": 19, "bold": True, "color": NAVY}])
    text(s, X0 + 7.05, yy, 1.25, 0.45, [{"runs": lab, "size": 19, "color": "595959"}])
    yy += 0.55
box(s, X0 + 8.45, 3.95, W_FULL - 8.45, 1.75, SAND, [
    {"runs": [(fmt_int(D['n_binary']), {"bold": True, "size": 28, "color": NAVY}), (" questions", {"size": 18})], "space_after": 2},
    {"runs": f"{D['classes']['mix_required']/D['n_binary']*100:.0f}% Mix required", "size": 17, "space_after": 2},
    {"runs": f"split {D['split_n']['train']} / {D['split_n']['validation']} / {D['split_n']['test']}, stratified by question type and label", "size": 15, "line": 1.0}],
    pad=0.22, anchor=MSO_ANCHOR.MIDDLE)
text(s, X0, 5.95, W_FULL, 0.7, [{"runs": "If both modes are correct, the judge is told to choose Naive: the label records the minimum sufficient mode, not the better-looking answer.",
                                  "size": 17, "color": DGREY}])
notes(s, f"""
[6:25 to 7:25]
The judge assigns one of three classes. If Naive's answer is correct, the question is naive_enough. If Naive is not enough but Mix is correct, it is mix_required. If neither is correct, none_enough.
Importantly, when both modes are correct the judge is instructed to choose Naive, so the label is the cheapest mode that was sufficient.
Out of {tot} judged questions, {D['classes']['naive_enough']} are naive_enough, {D['classes']['mix_required']} mix_required, and {D['classes']['none_enough']}, about a third, are none_enough. Those are excluded from training, because when both modes fail, choosing between them cannot improve the answer. I come back to what that third means at the end.
That leaves {D['n_binary']} questions, about a quarter of them Mix-required, split into train, validation and test, stratified by question type and label.
""")

# ============================================================================================ 10
s = content_slide("Two query-only routers: a lexical baseline and a fine-tuned encoder")
cw = 5.86
box(s, X0, 2.3, cw, 3.15, LIGHT, [
    {"runs": "TF-IDF + logistic regression", "bold": True, "size": 24, "color": NAVY, "space_after": 10},
    {"runs": "Word and word-pair counts, weighted by rarity", "size": 19, "bullet": True},
    {"runs": "Balanced class weights", "size": 19, "bullet": True},
    {"runs": "Trains in seconds on a CPU", "size": 19, "bullet": True},
    {"runs": "Tests how much of the signal is lexical", "size": 19, "bullet": True}], pad=0.3)
box(s, X0 + cw + 0.3, 2.3, cw, 3.15, KTH, [
    {"runs": "Fine-tuned ModernBERT", "bold": True, "size": 24, "color": WHITE, "space_after": 10},
    {"runs": "Pretrained encoder + classification head", "size": 19, "color": WHITE, "bullet": True},
    {"runs": "Class-balanced loss, early stopping", "size": 19, "color": WHITE, "bullet": True},
    {"runs": "3 seeds averaged into one ensemble", "size": 19, "color": WHITE, "bullet": True},
    {"runs": "Trains in hours on a CPU", "size": 19, "color": WHITE, "bullet": True}], pad=0.3)
box(s, X0, 5.75, W_FULL, 0.95, SAND, [
    {"runs": [("Both: ", {"bold": True, "color": NAVY}), ("threshold chosen on validation F1, then frozen.   ", {}),
              ("Reference: ", {"bold": True, "color": NAVY}), ("type-only rule (Mix if the question is compositional); not deployable, since real queries have no type label.", {})],
     "size": 18, "line": 1.05}], anchor=MSO_ANCHOR.MIDDLE, pad=0.25)
notes(s, """
[7:25 to 8:25]
I compare two routers that both see only the question text.
The baseline is TF-IDF with logistic regression: it represents a question by its words and word pairs, weighted by how rare they are. It trains in seconds and tells us how much of the routing signal is purely lexical.
The second is ModernBERT, a recent pretrained encoder, fine-tuned with a class-balanced loss and early stopping. I trained it with three random seeds and average their scores into one ensemble, so the result does not hinge on one lucky seed.
For both, the decision threshold is chosen on the validation split and then frozen before touching the test set.
As a reference I also include a rule that sends a question to Mix whenever the benchmark calls it compositional. It cannot be deployed, because real questions have no such label, but it shows how much of the signal is simply question structure.
""")

# ============================================================================================ 11
s = content_slide("Each router is evaluated as a classifier and as a routing policy, with paired tests")
cols = (("Classification", LIGHT, NAVY, (
            "AUPRC on mix_required (primary, chance 0.254)",
            "Precision, recall, F1",
            "Routing-label accuracy")),
        ("Routing policy", SAND, NAVY, (
            "Mean routed time (measured)",
            "Model tokens (estimated per mode)",
            "Under-routing: missed Mix call (answer risk)",
            "Over-routing: unneeded Mix call (wasted cost)")),
        ("Statistics", KTH, WHITE, (
            f"Same {D['n_test']} test questions for all routers, so tests are paired",
            "Paired bootstrap 95% intervals",
            "Permutation tests, McNemar",
            "Holm correction (7 metrics)")))
cw = (W_FULL - 0.6) / 3
for i, (head, fill, col, items) in enumerate(cols):
    paras = [{"runs": head, "bold": True, "size": 23, "color": col, "space_after": 12}]
    paras += [{"runs": it, "size": 18, "color": col if col == WHITE else DGREY, "bullet": True, "space_after": 8, "line": 1.0} for it in items]
    box(s, X0 + i * (cw + 0.3), 2.3, cw, 4.25, fill, paras, pad=0.28)
notes(s, f"""
[8:25 to 9:25]
Evaluation has two sides. As classifiers, the main metric is AUPRC on the minority class, mix_required, because the classes are imbalanced; a random router would score 0.254, the share of Mix-required questions. I also report precision, recall, F1 and accuracy.
As routing policies, I simulate what each router would have cost on the test set, using the precomputed answers and times: mean routed time, which is measured for every query, and model tokens, which are estimated from per-mode means. I separate the two kinds of mistakes: under-routing, sending a Mix-needing question to Naive, which risks the answer, and over-routing, sending a Naive-sufficient question to Mix, which only wastes cost.
Because all routers are scored on the same {D['n_test']} questions, the comparisons are paired: paired bootstrap intervals for differences, permutation tests, McNemar's test on individual routing decisions, and Holm correction for testing seven metrics.
""")

# ============================================================================================ 12
s = chapter_slide("Results", "Prediction (sub-question 1), then cost and risk (sub-question 2), then where the signal comes from")
notes(s, "[9:25]  Now the results: first how well the routers predict the label, which is sub-question one; then what routing costs and risks, which is sub-question two; then where the signal comes from.")

# ============================================================================================ 13
s = content_slide("Both routers find Mix-needing questions far above chance, and ModernBERT is more precise")
img(s, "c1_classification", X0, 2.15, 8.6, 4.55, "Bar chart comparing TF-IDF and ModernBERT on AUPRC, AUROC, precision, recall, F1 and accuracy on the 193-question test set; precision and accuracy differences significant.")
box(s, 9.55, 2.3, 3.13, 4.3, LIGHT, [
    {"runs": "Chance level", "bold": True, "size": 18, "color": NAVY, "space_after": 2},
    {"runs": "AUPRC 0.254", "size": 17, "space_after": 0},
    {"runs": "AUROC 0.5", "size": 17, "space_after": 14},
    {"runs": "Significant (Holm) *", "bold": True, "size": 18, "color": NAVY, "space_after": 2},
    {"runs": "Precision, accuracy", "size": 17, "space_after": 14},
    {"runs": "Recall is lower", "bold": True, "size": 18, "color": NAVY, "space_after": 2},
    {"runs": f"{CMP['recall_mix']['modernbert']:.2f} vs {CMP['recall_mix']['tfidf']:.2f}, not significant", "size": 17}], pad=0.25)
notes(s, f"""
[9:25 to 10:40]
Here are both routers on the {D['n_test']} test questions. TF-IDF is light blue, ModernBERT dark blue.
First the good news: both find Mix-needing questions far above chance. AUPRC is {CMP['auprc']['tfidf']:.2f} for TF-IDF and {CMP['auprc']['modernbert']:.2f} for ModernBERT, about three times the chance level of 0.254 shown by the dashed line, and AUROC is around 0.9. So the routing decision is partly predictable from the question alone. That is the first half of the answer to sub-question one.
ModernBERT has the higher point estimate on five of the six metrics. The clearest difference is precision: {CMP['precision_mix']['modernbert']:.2f} against {CMP['precision_mix']['tfidf']:.2f}. When ModernBERT says Mix, it is usually right.
The exception is recall: TF-IDF catches more of the Mix-needing questions, {CMP['recall_mix']['tfidf']:.2f} against {CMP['recall_mix']['modernbert']:.2f}, because it sends more questions to Mix.
Which of these differences are real, and which could be noise on 193 questions? That is the next slide.
""")

# ============================================================================================ 14
s = content_slide("Paired tests: the precision and accuracy gains are significant; the ranking gains are not resolved")
order = (("Precision", "precision_mix"), ("Accuracy", "accuracy"), ("F1", "f1_mix"), ("AUPRC", "auprc"), ("AUROC", "auroc"),
         ("Balanced accuracy", "balanced_accuracy"), ("Recall", "recall_mix"))
lo_v, hi_v = -0.25, 0.35
px0, px1 = X0 + 2.55, X0 + 8.3
def vx(v):
    return px0 + (v - lo_v) / (hi_v - lo_v) * (px1 - px0)
ytop, rh = 2.5, 0.455
text(s, X0 + 8.5, ytop - 0.45, 3.6, 0.35, [{"runs": "difference  [95% interval]       Holm p", "size": 14, "color": "595959"}])
for i, (name, k) in enumerate(order):
    c = CMP[k]
    yy = ytop + i * rh
    signif = c["perm_p_holm"] < 0.05
    col = KTH if signif else GREY
    text(s, X0, yy - 0.02, 2.45, 0.42, [{"runs": name, "size": 18, "bold": signif, "color": NAVY if signif else DGREY}])
    line(s, vx(c["diff_ci95"][0]), yy + 0.2, vx(c["diff_ci95"][1]), yy + 0.2, col, 3.5 if signif else 2.5)
    dot(s, vx(c["diff"]), yy + 0.2, 0.22, col if signif else WHITE, None if signif else GREY)
    lo, hi = c["diff_ci95"]
    text(s, X0 + 8.5, yy - 0.02, 2.6, 0.42, [{"runs": f"{sgn(c['diff'])}  [{sgn(lo)}, {sgn(hi)}]", "size": 15, "bold": signif, "color": NAVY if signif else DGREY}])
    ph = c["perm_p_holm"]
    text(s, X0 + 11.1, yy - 0.02, 0.95, 0.42, [{"runs": "<0.001" if ph < 0.001 else f"{ph:.3f}" if ph < 0.2 else f"{ph:.2f}", "size": 15, "bold": signif, "color": NAVY if signif else DGREY}])
yb = ytop + len(order) * rh + 0.05
line(s, vx(0), ytop - 0.15, vx(0), yb, DGREY, 1.25, dash="dash")
line(s, px0, yb, px1, yb, GREY, 1)
for t_ in (-0.2, -0.1, 0.0, 0.1, 0.2, 0.3):
    text(s, vx(t_) - 0.4, yb + 0.05, 0.8, 0.3, [{"runs": sgn(t_, 1) if t_ else "0", "align": PP_ALIGN.CENTER, "size": 13, "color": "595959"}])
text(s, px0 - 0.4, yb + 0.33, px1 - px0 + 0.8, 0.3, [{"runs": f"ModernBERT {MINUS} TF-IDF  (right of zero favours ModernBERT)", "align": PP_ALIGN.CENTER, "size": 14, "color": "595959"}])
box(s, X0, 6.38, W_FULL, 0.58, SAND, [{"runs": [("Reading: ", {"bold": True, "color": NAVY}),
    (f"ModernBERT is the more selective router. Whether it also ranks questions better cannot be decided with {D['n_pos']} Mix-required test questions.", {})], "size": 17}],
    anchor=MSO_ANCHOR.MIDDLE, pad=0.2)
notes(s, f"""
[10:40 to 12:10]
This slide is the statistical core, so let me say how to read it. Each row is one metric. The dot is the difference, ModernBERT minus TF-IDF, and the bar is its 95 percent paired-bootstrap interval. The dashed line is zero: if a bar crosses it, the data cannot tell the two routers apart on that metric. Blue rows are significant after Holm correction for testing seven metrics; grey rows are not.
Precision: plus {CMP['precision_mix']['diff']:.2f}, with the whole interval well to the right of zero, Holm-adjusted p below 0.001. Accuracy: plus {CMP['accuracy']['diff']:.3f}, adjusted p {CMP['accuracy']['perm_p_holm']:.3f}. McNemar's test agrees: ModernBERT gets {MC['mb_vs_tfidf_all']['a_right_b_wrong']} questions right that TF-IDF gets wrong, and the reverse happens {MC['mb_vs_tfidf_all']['a_wrong_b_right']} times.
Everything else, F1, AUPRC, AUROC, balanced accuracy, and the lower recall, has an interval that includes zero. I read that as undecided, not as equal: with only {D['n_pos']} Mix-required test questions, differences of a few hundredths in AUPRC cannot be resolved.
So the honest answer to sub-question one is: both routers predict the label well above chance, and ModernBERT is significantly more selective than TF-IDF. I cannot claim it ranks questions better overall.
""")

# ============================================================================================ 15
s = content_slide("ModernBERT makes a third as many unnecessary Mix calls, at a small cost in missed ones")
img(s, "c2_errors", X0, 2.15, 7.6, 4.6, "Bar chart of routing errors: unnecessary Mix calls TF-IDF 30 vs ModernBERT 10; missed Mix calls TF-IDF 11 vs ModernBERT 16.")
mo, mu = MC["mb_vs_tfidf_naive_rows(over_routing)"], MC["mb_vs_tfidf_mix_rows(recall)"]
box(s, 8.6, 2.3, 4.08, 2.05, LIGHT, [
    {"runs": "Over-routing: wasted cost", "bold": True, "size": 19, "color": NAVY, "space_after": 6},
    {"runs": f"{tf_['fp']} → {mb['fp']} unnecessary Mix calls", "size": 18, "space_after": 4},
    {"runs": f"McNemar: {mo['a_right_b_wrong']} vs {mo['a_wrong_b_right']}, {pnb(mo['p_value'])}", "size": 17, "bold": True, "color": KTH}], pad=0.22)
box(s, 8.6, 4.5, 4.08, 2.05, SAND, [
    {"runs": "Under-routing: answer risk", "bold": True, "size": 19, "color": NAVY, "space_after": 6},
    {"runs": f"{tf_['fn']} → {mb['fn']} missed Mix calls", "size": 18, "space_after": 4},
    {"runs": f"McNemar: {mu['a_wrong_b_right']} vs {mu['a_right_b_wrong']}, {pnb(mu['p_value'])} (not significant)", "size": 17}], pad=0.22)
notes(s, f"""
[12:10 to 13:10]
Translating that into routing decisions. On the left, unnecessary Mix calls: questions where Naive was enough but the router chose Mix. TF-IDF makes {tf_['fp']} of them, ModernBERT {mb['fp']}. That difference is clear: ModernBERT avoids {mo['a_right_b_wrong']} of TF-IDF's mistakes and makes no new ones, McNemar p well below 0.001.
On the right, missed Mix calls, which are the dangerous error because the cheap mode then fails. ModernBERT misses {mb['fn']}, TF-IDF {tf_['fn']}. That gap rests on {mu['a_wrong_b_right']} versus {mu['a_right_b_wrong']} questions and is not significant.
So the trade is: far less wasted compute, for a small and statistically unresolved increase in risk. Which side you prefer depends on what a wasted Mix call costs compared with a degraded answer in a given deployment.
""")

# ============================================================================================ 16
s = content_slide(f"ModernBERT routing cuts mean time by {mb['time_saving']*100:.0f}% and estimated tokens by {mb['token_saving']*100:.0f}% compared with always-Mix")
img(s, "c3_cost", X0, 2.2, W_FULL, 3.95, "Two bar charts: mean time per query and estimated model tokens per query for always-Naive, type-only rule, TF-IDF, ModernBERT and always-Mix.")
rt = D["routed_time_test"]
box(s, X0, 6.22, W_FULL, 0.72, LIGHT, [
    {"runs": [("ModernBERT vs TF-IDF time: ", {"bold": True, "color": NAVY}),
              (f"{sgn(rt['diff_mean_seconds'], 1)} s per query, 95% interval [{num(rt['diff_ci95'][0], 1)}, {num(rt['diff_ci95'][1], 1)}], {pnb(rt['perm_p'])}.", {})], "size": 16, "space_after": 0},
    {"runs": f"*Tokens: per-mode means from {TM['n']} instrumented queries, no interval. The router's own inference time is not included in either chart.", "size": 14, "color": "595959"}],
    anchor=MSO_ANCHOR.MIDDLE, pad=0.18)
notes(s, f"""
[13:10 to 14:25]
Now sub-question two, cost. Left: mean time per query, measured on every test question. Always-Naive takes {an['mean_time']:.1f} seconds, always-Mix {am['mean_time']:.1f}. TF-IDF routing gets {tf_['mean_time']:.1f}, ModernBERT {mb['mean_time']:.1f}: a {mb['time_saving']*100:.0f} percent reduction relative to always-Mix. And the difference between the two routers, about {abs(rt['diff_mean_seconds']):.1f} seconds per query, is itself significant, p = {rt['perm_p']:.3f}.
Right: model tokens per query. ModernBERT routes at about {fmt_int(TOK['modernbert'])} tokens against {fmt_int(TOK['always_mix'])} for always-Mix, a {mb['token_saving']*100:.0f} percent reduction; TF-IDF saves {tf_['token_saving']*100:.0f} percent. The saving is larger in tokens than in time because Mix adds an extra language-model call and a much larger context, while wall-clock time also includes network latency.
Two caveats I want to be upfront about. The token figures are estimates, from per-mode means measured on {TM['n']} queries; the time figures are measured for every query. And the router's own inference time, three encoder passes for the ensemble, is not included. I expect it to be small next to retrieval calls of 11 to 24 seconds, but I did not measure it.
""")

# ============================================================================================ 17
s = content_slide(f"ModernBERT keeps about {mb['retention']*100:.0f}% of judged-correct answers at {TOK['modernbert']/TOK['always_mix']*100:.0f}% of always-Mix's tokens")
img(s, "c4_tradeoff", X0, 2.15, 8.3, 4.65, "Scatter of estimated tokens per query against share of always-Mix's judged-correct answers kept, for five routing policies.")
box(s, 9.2, 2.3, 3.48, 2.5, LIGHT, [
    {"runs": "ModernBERT", "bold": True, "size": 20, "color": NAVY, "space_after": 4},
    {"runs": f"keeps {mb['retention']*100:.1f}% of the answers at {TOK['modernbert']/TOK['always_mix']*100:.0f}% of always-Mix's token cost", "size": 18, "space_after": 10},
    {"runs": "TF-IDF", "bold": True, "size": 20, "color": NAVY, "space_after": 4},
    {"runs": f"keeps {tf_['retention']*100:.1f}% at {TOK['tfidf']/TOK['always_mix']*100:.0f}%", "size": 18}], pad=0.22)
box(s, 9.2, 5.05, 3.48, 1.75, SAND, [{"runs": "Assumes Mix is correct wherever Naive was, so kept answers = 1 − under-routing. Answers were not re-judged per policy.",
                                       "size": 15, "line": 1.05}], pad=0.2)
notes(s, f"""
[14:25 to 15:25]
What does the saving cost in answers? I did not re-judge the answers each policy would pick, so this uses one explicit assumption: that Mix is correct whenever Naive was. Then the only way a policy can lose an answer is by under-routing, and the share of always-Mix's correct answers it keeps is one minus the under-routing rate.
Plotted against estimated tokens: always-Naive is cheapest but keeps only {an['retention']*100:.1f} percent. ModernBERT keeps {mb['retention']*100:.1f} percent at {TOK['modernbert']/TOK['always_mix']*100:.0f} percent of always-Mix's token cost. TF-IDF and the type-only rule keep {tf_['retention']*100:.1f} percent at about half the cost.
So ModernBERT gives up about eight points of answers, {mb['retention']*100:.1f} against 100, in exchange for cutting almost two thirds of the estimated token cost. And remember from slide 15 that the difference between the two routers in missed Mix calls is not significant, whereas the difference in time is.
""")

# ============================================================================================ 18
s = content_slide("The threshold is a cost dial, and both chosen operating points sit well above the static policies")
img(s, "c5_frontier", X0, 2.15, 8.6, 4.65, "Cost-performance frontier: F1 on mix_required against mean routed time for TF-IDF and ModernBERT as the threshold sweeps, with the two static policies.")
box(s, 9.55, 2.3, 3.13, 4.4, LIGHT, [
    {"runs": "Reading the curve", "bold": True, "size": 19, "color": NAVY, "space_after": 8},
    {"runs": "Each point is one threshold; lowering it sends more questions to Mix.", "size": 17, "space_after": 10, "line": 1.0},
    {"runs": "ModernBERT's chosen threshold sits near its peak; TF-IDF's is a few hundredths below its own (0.65 vs 0.69).", "size": 17, "space_after": 10, "line": 1.0},
    {"runs": "In deployment the threshold sets the cost.", "size": 17, "bold": True, "color": NAVY, "line": 1.0}], pad=0.22)
notes(s, f"""
[15:25 to 16:25]  (shorten to one sentence if running long)
Routing is not one fixed point. Each router has a threshold, and sweeping it from 0 to 1 traces a curve from always-Mix on the right to always-Naive on the left. The y-axis is F1 on the Mix-required class, the x-axis mean routed time.
Over a broad middle range both routers sit well above always-Mix: higher routing F1 at much lower time. At the low-time end, below about 13 seconds, F1 falls under always-Mix's, because the router then sends almost everything to Naive.
The diamonds are the thresholds chosen on validation. ModernBERT's, {D['thr']['modernbert']:.2f}, sits near its peak. TF-IDF's, {D['thr']['tfidf']:.2f}, is a few hundredths below the best point on its own curve. The plateau is broad, so the choice is not brittle.
Practically, the threshold is a dial a deployment can turn to match its own cost of a wasted Mix call against a missed one.
""")

# ============================================================================================ 19
s = content_slide("Much of the signal is question structure: compositional questions are the ones that usually need Mix")
img(s, "c6_types", X0, 2.65, 6.2, 4.0, "Bar chart of the share of Mix-required questions per 2WikiMultihopQA question type in the binary dataset; compositional 58 percent.")
text(s, X0, 2.2, 6.2, 0.4, [{"runs": f"Mix-required rate by question type (all {fmt_int(D['n_binary'])} labelled questions)", "size": 17, "bold": True, "color": NAVY}])
vt = D["vs_type_mb"]; tp = D["type_only_point"]; pc = D["per_type_tests"]["compositional"]["precision_mix"]
vt_hp = {k: v.get("p_value_holm", v["p_value"]) for k, v in vt.items()}
mb_bridge = D["per_type_metrics"]["modernbert_ensemble"]["bridge_comparison"]
cx, cw = 7.25, 5.43
cards = (
    (SAND, "Type-only rule", f"F1 {tp['f1_mix']:.2f}, using the benchmark's type label, which real queries do not have"),
    (LIGHT, "ModernBERT vs the rule (text only)", f"Precision {vt['precision_mix']['diff']+tp['precision_mix']:.2f} vs {tp['precision_mix']:.2f} (Holm {pnb(vt_hp['precision_mix'])}); accuracy higher only before correction"),
    (LIGHT, "Within compositional questions", f"ModernBERT precision {pc['modernbert']:.2f} vs TF-IDF {pc['tfidf']:.2f} ({pnb(pc['perm_p'])}, exploratory)"),
    (LBRICK, "Weak spot: bridge-comparison", f"ModernBERT finds {round(mb_bridge['recall_mix_required']*mb_bridge['n_mix'])} of {mb_bridge['n_mix']} Mix-required questions (backup 26)"))
yy = 2.25
for fill, head, body in cards:
    box(s, cx, yy, cw, 1.05, fill, [{"runs": head, "bold": True, "size": 18, "color": NAVY, "space_after": 2},
                                    {"runs": body, "size": 15, "line": 1.0}], pad=0.18, anchor=MSO_ANCHOR.MIDDLE)
    yy += 1.15
notes(s, f"""
[16:25 to 17:40]
Where does the signal come from? This chart shows how often each benchmark question type needs Mix. Comparison questions almost never do, {bt['comparison']['mix']} of {bt['comparison']['n']}. Compositional questions, like the Private Hell 36 example, need it {bt['compositional']['mix']/bt['compositional']['n']*100:.0f} percent of the time, and they make up {bt['compositional']['mix']} of the {D['classes']['mix_required']} Mix-required questions.
That is why the type-only rule is strong: Mix for compositional, Naive otherwise, reaches F1 {tp['f1_mix']:.2f}. But it uses a label that real queries do not have.
ModernBERT, reading only the text, is significantly more precise than that rule. Its accuracy advantage holds only before multiple-testing correction, and TF-IDF does not beat the rule on anything. Within compositional questions, ModernBERT's precision advantage over TF-IDF is also significant, as an exploratory test.
And there is a clear weak spot: on bridge-comparison questions ModernBERT finds only {round(mb_bridge['recall_mix_required']*mb_bridge['n_mix'])} of {mb_bridge['n_mix']} Mix-required ones.
So I read the router as learning question structure from phrasing, plus some precision on top. Useful, but part of the signal is benchmark-specific, which matters for transfer.
""")

# ============================================================================================ 20
s = content_slide("What this means for a deployment like Ericsson's")
cards = (("Cheap to run, tunable", LIGHT, NAVY, "No language-model call at query time; the threshold sets the cost. Trains in hours on a CPU. Labels cost one run of both modes plus the judge per training question."),
         ("Validate on the target workload", SAND, NAVY, "Part of the signal is benchmark structure. Real queries need their own labels: a small human-labelled set or a judge that does not rely on gold answers."),
         ("Routing cannot fix everything", KTH, WHITE, f"{D['classes']['none_enough']/D['n_judged']*100:.0f}% of the questions failed in both modes. They need a better answering pipeline or a fallback, not a better router."))
cw = (W_FULL - 0.6) / 3
for i, (h, fill, col, body) in enumerate(cards):
    box(s, X0 + i * (cw + 0.3), 2.3, cw, 4.35, fill, [{"runs": h, "bold": True, "size": 22, "color": col, "space_after": 14, "line": 1.0},
                                                       {"runs": body, "size": 18, "color": col if col == WHITE else DGREY, "line": 1.08}], pad=0.3)
notes(s, f"""
[17:40 to 18:40]  (shorten to one sentence if running long)
What does this mean in practice, for example at Ericsson where the corpus is internal technical documentation?
First, the router is cheap to run and tunable: no language-model call at query time, and the threshold sets where you sit on the cost curve. It is not free to set up, though: every training label requires running both modes and the judge once.
Second, it must be validated on the target workload. Part of what it learned is the structure of this benchmark. Internal documentation has no gold answers, so labelling would need a small human-labelled set or a different judging protocol.
Third, routing has a ceiling: {D['classes']['none_enough']/D['n_judged']*100:.0f} percent of the questions failed in both modes. For those, choosing between Naive and Mix cannot help; they need a better end-to-end pipeline.
One ethical point: a router can serve some kinds of question systematically worse. Here ModernBERT sent {mb_bridge['n_mix'] - round(mb_bridge['recall_mix_required']*mb_bridge['n_mix'])} of the {mb_bridge['n_mix']} Mix-required bridge-comparison questions to Naive, so a deployment should monitor quality per question segment, not only on average. On sustainability: fewer Mix calls mean fewer tokens and less time, which should lower energy use, although I did not measure energy.
""")

# ============================================================================================ 21
s = content_slide("Limitations, and what would remove them")
lim = (("Labels come from an LLM judge, using the same Gemini model that wrote the answers; no human check", "Hand-validate a sample; re-judge with a second model family"),
       (f"Token cost estimated from {TM['n']} instrumented queries", "Log tokens for every query"),
       ("One benchmark (chosen partly for label balance), one framework, row-random split", "Template-grouped split; HotpotQA, MuSiQue; other frameworks"),
       (f"{D['n_test']} test questions ({D['n_pos']} Mix-required) cannot resolve small ranking differences", "A test set 5 to 20 times larger"),
       ("Kept answers assume Mix is correct wherever Naive was", "Re-judge the answer each policy picks"))
text(s, X0, 2.15, 6.6, 0.4, [{"runs": "Limitation", "bold": True, "size": 17, "color": KTH}])
text(s, X0 + 7.0, 2.15, 5.0, 0.4, [{"runs": "Next step", "bold": True, "size": 17, "color": KTH}])
yy = 2.6
for a, b in lim:
    text(s, X0, yy, 6.6, 0.8, [{"runs": a, "size": 18, "line": 1.0}], anchor=MSO_ANCHOR.MIDDLE)
    arrow(s, X0 + 6.65, yy + 0.4, X0 + 6.95, yy + 0.4, GREY, 1.5)
    box(s, X0 + 7.0, yy + 0.04, W_FULL - 7.0, 0.72, LIGHT, [{"runs": b, "size": 17, "line": 1.0, "color": NAVY}], anchor=MSO_ANCHOR.MIDDLE, pad=0.15)
    yy += 0.86
notes(s, """
[18:40 to 19:40]
I want to be explicit about the limits.
The labels come from an LLM judge, and the same Gemini model both wrote and judged the answers, with no human check. Giving the judge the gold answer reduces the risk, but shared bias is possible; the fix is a hand-validated sample and a second judge from another model family.
Token costs come from 15 instrumented queries, so they are preliminary; the fix is logging tokens for every query.
It is one benchmark, chosen partly because it gave a balanced label split, one framework, and a split that is random by row, so similar question templates appear on both sides.
193 test questions with 49 positives cannot resolve small ranking differences.
And the kept-answers figure rests on an assumption that re-judging each policy's answers would remove.
What these do not change: the time saving is measured on every test query, and the precision gain is significant on the labels as defined. What they limit is how far the numbers transfer.
""")

# ============================================================================================ 22
s = prs.slides.add_slide(LAYOUT["Comparison, KTH blue"])
s.shapes.title.text_frame.text = "Conclusions"
vt_prec_rule = vt["precision_mix"]["diff"] + tp["precision_mix"]
for ph in s.placeholders:
    idx = ph.placeholder_format.idx
    if idx == 1:
        tf = ph.text_frame; tf.clear()
        set_runs(tf, [
            {"runs": [("Prediction (sub-question 1). ", {"bold": True, "color": NAVY}), (f"From the question alone, both routers find Mix-needing questions at about 3× chance. ModernBERT is significantly more precise than TF-IDF ({CMP['precision_mix']['modernbert']:.2f} vs {CMP['precision_mix']['tfidf']:.2f}) and than the type-only rule; ranking gains are unresolved.", {})], "size": 17, "space_after": 12, "line": 1.05},
            {"runs": [("Cost (sub-question 2). ", {"bold": True, "color": NAVY}), (f"ModernBERT routing cuts mean time by {mb['time_saving']*100:.0f}% and estimated tokens by {mb['token_saving']*100:.0f}% versus always-Mix, keeping about {mb['retention']*100:.0f}% of judged-correct answers (assuming Mix is correct wherever Naive was).", {})], "size": 17, "space_after": 12, "line": 1.05},
            {"runs": [("Scope. ", {"bold": True, "color": NAVY}), ("One benchmark, one framework, LLM-judged labels: validate on the target workload before deployment.", {})], "size": 17, "line": 1.05}], size=17)
    elif idx == 14:
        ph.text_frame.text = "Contributions"
    elif idx == 13:
        tf = ph.text_frame; tf.clear()
        set_runs(tf, [
            {"runs": "A study of query-only routing between a vector and a graph-enhanced retrieval mode within one RAG framework", "size": 18, "color": WHITE, "space_after": 12, "bullet": True, "line": 1.05},
            {"runs": "A judge-labelled routing dataset and an open, reproducible pipeline", "size": 18, "color": WHITE, "space_after": 12, "bullet": True, "line": 1.05},
            {"runs": "Evidence, from paired tests, that the encoder's gain lies in precision and therefore in cost", "size": 18, "color": WHITE, "bullet": True, "line": 1.05}], size=18)
text(s, 7.29, 5.75, 5.39, 1.0, [{"runs": "Thank you. Questions are welcome.", "size": 22, "bold": True, "color": WHITE, "space_after": 4},
                                 {"runs": "github.com/Dantecool02/Thesis_ericsson", "size": 14, "color": LIGHT}])
add_footer(s, footer=False)
notes(s, f"""
[19:40 to 20:40]  Leave this slide up during the questions.
To conclude, and to answer the two sub-questions directly.
One, prediction: from the question alone, both routers find Mix-needing questions at about three times chance. ModernBERT is significantly more precise than TF-IDF and than the type-only rule; whether it ranks questions better is unresolved at this test size.
Two, cost: routing with ModernBERT cuts mean time by {mb['time_saving']*100:.0f} percent and estimated tokens by {mb['token_saving']*100:.0f} percent compared with always running Mix, while keeping about {mb['retention']*100:.0f} percent of judged-correct answers under the stated assumption.
Concretely, my part of the work was to design the labelling protocol, run both modes on all 2,000 questions, build the dataset, train both routers and do the statistical analysis; Ericsson provided the problem setting and supervision. [Check this sentence is accurate before the talk.]
The contributions are the study itself, a judge-labelled routing dataset with an open and reproducible pipeline, and the evidence that the encoder's advantage lies in precision, which is exactly what translates into cost.
Thank you, I am happy to take questions.
""")

# ============================================================================================ BACKUP
s = chapter_slide("Backup slides", "Type the slide number in Keynote and press Return to jump", layout="Title/chapter, navy blue")
idx_items = (("24", "Full statistics table"), ("25", "Seed stability"), ("26", "Per question type"), ("27", "Cost and errors per policy"),
             ("28", "Dataset composition"), ("29", "Judge prompt"), ("30", "PR and ROC curves"), ("31", "Statistical power"),
             ("32", "Query-only vs Naive-first cascade"), ("33", "Configuration"), ("34", "Use of generative AI"))
text(s, 8.65, 1.25, 4.2, 5.0, [{"runs": [(f"{n_}  ", {"bold": True, "color": WHITE}), (lab, {"color": LIGHT})], "size": 17, "space_after": 6} for n_, lab in idx_items])
notes(s, "Backup section. Jump here by slide number when a question needs the detail (type the number and press Return in Keynote).")

def table(slide, x, y, w, h, data, col_w=None, size=15, header_fill=KTH, zebra=LIGHT, bold_rows=(), first_col_bold=False, left_cols=(0,)):
    rows_n, cols_n = len(data), len(data[0])
    gt = slide.shapes.add_table(rows_n, cols_n, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gt.table
    sid = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"
    if col_w:
        for i, cw_ in enumerate(col_w):
            tbl.columns[i].width = Inches(cw_)
    for r in range(rows_n):
        tbl.rows[r].height = Inches(h / rows_n)
        for c in range(cols_n):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(header_fill if r == 0 else (zebra if r % 2 == 0 else WHITE))
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.text = ""
            run = p.add_run(); run.text = str(data[r][c])
            run.font.size = Pt(size)
            run.font.bold = r == 0 or r in bold_rows or (first_col_bold and c == 0)
            run.font.color.rgb = rgb(WHITE if r == 0 else DGREY)
            p.alignment = PP_ALIGN.LEFT if c in left_cols else PP_ALIGN.CENTER
    return gt

def ci(d):
    return f"[{num(d['lo'])}, {num(d['hi'])}]"

# B1 (24) full paired table
s = content_slide("Backup: full classification comparison with intervals and paired tests")
mg = D["marginal"]; mp = D["marginal_pr"]
rows_ = [["Metric", "TF-IDF  [95% CI]", "ModernBERT  [95% CI]", "Difference  [95% paired CI]", "p", "Holm p"]]
spec = (("AUPRC", "auprc", mg["tfidf_baseline"]["auprc"], mg["modernbert_ensemble"]["auprc"]),
        ("AUROC", "auroc", mg["tfidf_baseline"]["auroc"], mg["modernbert_ensemble"]["auroc"]),
        ("Balanced accuracy", "balanced_accuracy", mg["tfidf_baseline"]["bal"], mg["modernbert_ensemble"]["bal"]),
        ("Precision (mix)", "precision_mix", {"lo": mp["tfidf"]["precision_mix_ci95"][0], "hi": mp["tfidf"]["precision_mix_ci95"][1]}, {"lo": mp["modernbert_ensemble"]["precision_mix_ci95"][0], "hi": mp["modernbert_ensemble"]["precision_mix_ci95"][1]}),
        ("Recall (mix)", "recall_mix", {"lo": mp["tfidf"]["recall_mix_ci95"][0], "hi": mp["tfidf"]["recall_mix_ci95"][1]}, {"lo": mp["modernbert_ensemble"]["recall_mix_ci95"][0], "hi": mp["modernbert_ensemble"]["recall_mix_ci95"][1]}),
        ("F1 (mix)", "f1_mix", mg["tfidf_baseline"]["f1"], mg["modernbert_ensemble"]["f1"]),
        ("Accuracy", "accuracy", mg["tfidf_baseline"]["acc"], mg["modernbert_ensemble"]["acc"]))
for name, k, a, b in spec:
    c = CMP[k]
    rows_.append([name, f"{c['tfidf']:.3f}  {ci(a)}", f"{c['modernbert']:.3f}  {ci(b)}",
                  f"{sgn(c['diff'])}  [{sgn(c['diff_ci95'][0])}, {sgn(c['diff_ci95'][1])}]",
                  "<0.001" if c["perm_p"] < 0.001 else f"{c['perm_p']:.3f}", "<0.001" if c["perm_p_holm"] < 0.001 else f"{c['perm_p_holm']:.3f}"])
table(s, X0, 2.25, W_FULL, 3.9, rows_, col_w=[2.25, 2.35, 2.55, 2.97, 0.95, 0.95], size=15, bold_rows=(4, 7))
text(s, X0, 6.3, W_FULL, 0.6, [{"runs": f"n = {D['n_test']} test questions ({D['n_pos']} mix_required). Marginal CIs: 1,000 bootstrap resamples. Paired CIs: 10,000 paired resamples. p: paired permutation test (10,000). Holm over the 7 metrics. DeLong AUROC p = {D['delong']['p_value']:.2f}.",
                                 "size": 13, "color": "595959"}])
notes(s, "Use for: 'Is ModernBERT actually better?' / 'Show me the statistics.' Thesis Table 4.2.")

# B2 (25) seeds
s = content_slide("Backup: every ModernBERT seed scores above TF-IDF on the aggregate metrics")
ps = D["per_seed"]; ss = D["seed_stability"]
rows_ = [["Run", "Threshold", "AUPRC", "AUROC", "Bal. acc.", "Precision", "Recall", "F1", "Accuracy"]]
keys = ("auprc", "auroc", "balanced_accuracy", "precision_mix", "recall_mix", "f1_mix", "accuracy")
for sd in ("7", "13", "42"):
    rows_.append([f"Seed {sd}", f"{ps[sd]['threshold']:.2f}"] + [f"{ps[sd][k]:.3f}" for k in keys])
rows_.append(["Sample std", "–"] + [f"{ss[k]['std_ddof1']:.3f}" for k in keys])
rows_.append(["Ensemble", f"{D['thr']['modernbert']:.2f}"] + [f"{CMP[k]['modernbert']:.3f}" for k in keys])
rows_.append(["TF-IDF", f"{D['thr']['tfidf']:.2f}"] + [f"{CMP[k]['tfidf']:.3f}" for k in keys])
table(s, X0, 2.3, W_FULL, 3.5, rows_, col_w=[1.7, 1.3] + [1.29] * 7, size=16, bold_rows=(5,))
text(s, X0, 6.0, W_FULL, 0.8, [{"runs": "Point estimates. Aggregate metrics vary little across seeds; precision and recall vary more because each seed selects its own validation threshold. The ensemble costs three encoder forward passes per query.",
                                 "size": 16, "color": DGREY}])
notes(s, "Use for: 'Why an ensemble?' / 'Is it a lucky seed?' Thesis Table 4.5.")

# B3 (26) per type
s = content_slide("Backup: per-question-type results on the test set")
pt_ = D["per_type_metrics"]
rows_ = [["Slice", "n", "n mix", "Router", "AUPRC", "Precision", "Recall", "F1", "Accuracy"]]
for k, n_ in (("comparison", "Comparison"), ("compositional", "Compositional"), ("bridge_comparison", "Bridge-comparison"), ("inference", "Inference")):
    for j, (m, mn) in enumerate((("tfidf_baseline", "TF-IDF"), ("modernbert_ensemble", "ModernBERT"))):
        r = pt_[m][k]
        rows_.append([n_ if j == 0 else "", r["n"] if j == 0 else "", r["n_mix"] if j == 0 else "", mn, f"{r['auprc']:.3f}",
                      f"{r['precision_mix_required']:.3f}", f"{r['recall_mix_required']:.3f}", f"{r['f1_mix_required']:.3f}", f"{r['accuracy']:.3f}"])
table(s, X0, 2.2, W_FULL, 4.0, rows_, col_w=[2.3, 0.7, 0.9, 1.7, 1.28, 1.28, 1.28, 1.28, 1.3], size=15)
pcs = D["per_type_tests"]["compositional"]
text(s, X0, 6.3, W_FULL, 0.6, [{"runs": f"Exploratory, unadjusted tests on compositional and bridge-comparison only: just compositional precision differs significantly ({pnb(pcs['precision_mix']['perm_p'])}). Comparison and inference have too few Mix-required questions to test.",
                                 "size": 14, "color": "595959"}])
notes(s, "Use for: 'Is the router just learning the question type?' / per-slice questions / bridge-comparison weakness. Thesis Table 4.4.")

# B4 (27) cost table
s = content_slide("Backup: routed cost and routing errors per policy")
rows_ = [["Policy", "To Mix", "Mean time (s)", "Est. tokens", "Token saving", "Missed Mix", "Extra Mix", "Kept answers*"]]
for k in ("always_naive", "type_only", "tfidf", "modernbert", "always_mix"):
    p = P[k]
    rows_.append([p["name"], f"{p['mix_rate']*100:.1f}%", f"{p['mean_time']:.2f}", fmt_int(TOK[k]), f"{p['token_saving']*100:.1f}%",
                  f"{p['fn']} ({p['under']:.3f})", f"{p['fp']} ({p['over']:.3f})", f"{p['retention']*100:.1f}%"])
table(s, X0, 2.3, W_FULL, 3.2, rows_, col_w=[2.1, 1.1, 1.55, 1.45, 1.45, 1.55, 1.55, 1.27], size=16, bold_rows=(4,))
text(s, X0, 5.75, W_FULL, 1.0, [{"runs": f"Time: measured per query. Tokens: per-mode means from {TM['n']} instrumented queries ({fmt_int(TM['naive'])} Naive, {fmt_int(TM['mix'])} Mix; prompt, output and reasoning tokens; embeddings excluded). Missed Mix = under-routing, Extra Mix = over-routing. *Kept answers = 1 − under-routing, assuming Mix is correct wherever Naive was.",
                                 "size": 15, "color": DGREY}])
notes(s, "Use for: token-estimate questions, under/over-routing, answer-retention assumption. Thesis Table 4.3.")

# B5 (28) dataset
s = content_slide("Backup: dataset composition by split and question type")
stp = D["split_type"]
rows_ = [["Split", "Total", "Comparison", "Compositional", "Bridge-comparison", "Inference"]]
for sp, nm in (("train", "Train"), ("validation", "Validation"), ("test", "Test")):
    rows_.append([nm, D["split_n"][sp]] + [f"{stp[sp][t][0]} ({stp[sp][t][1]} mix)" for t in ("comparison", "compositional", "bridge_comparison", "inference")])
rows_.append(["Overall", fmt_int(D["n_binary"])] + [f"{bt[t]['n']} ({bt[t]['mix']} mix)" for t in ("comparison", "compositional", "bridge_comparison", "inference")])
table(s, X0, 2.3, W_FULL, 2.6, rows_, col_w=[1.8, 1.2, 2.2, 2.3, 2.42, 2.1], size=17, bold_rows=(4,))
text(s, X0, 5.2, W_FULL, 1.4, [
    {"runs": f"Label flow: {fmt_int(D['n_judged'])} judged → {D['classes']['naive_enough']} naive_enough, {D['classes']['mix_required']} mix_required, {D['classes']['none_enough']} none_enough (excluded) → {fmt_int(D['n_binary'])} binary.", "size": 17, "space_after": 6},
    {"runs": "Query set: first 2,000 of the 2,279 queryable questions of the 2WikiMultihopQA train split, in benchmark order. Split jointly stratified by question type and label.", "size": 17}])
notes(s, "Use for: 'Why exclude none_enough?' / 'How were questions selected?' Thesis Tables 3.1-3.2.")

# B6 (29) judge prompt
s = content_slide("Backup: the judge sees the gold answer and must return the minimum sufficient label")
box(s, X0, 2.25, W_FULL, 4.4, LIGHT, [
    {"runs": "Labels defined in the prompt (quoted)", "bold": True, "size": 18, "color": NAVY, "space_after": 4},
    {"runs": "naive_enough: the Naive answer was already enough to answer the question correctly.", "size": 16, "bullet": True, "space_after": 2},
    {"runs": "mix_required: the Naive answer was not enough, but the Mix answer provided more context that led to a correct answer.", "size": 16, "bullet": True, "space_after": 2},
    {"runs": "none_enough: neither answer was enough.", "size": 16, "bullet": True, "space_after": 12},
    {"runs": "Key rules (9 in total)", "bold": True, "size": 18, "color": NAVY, "space_after": 4},
    {"runs": "Use the ground-truth answer to judge correctness when it is provided. A longer answer is not better just because it is longer.", "size": 16, "bullet": True, "space_after": 2},
    {"runs": "If both Naive and Mix are correct, choose naive_enough.", "size": 16, "bullet": True, "space_after": 2},
    {"runs": "Treat explicit admissions of missing information and obvious error messages as not enough.", "size": 16, "bullet": True, "space_after": 12},
    {"runs": "Model gemini-3-flash-preview, temperature 0, one fixed prompt for all 2,000 questions. Full system and user prompt verbatim in thesis Appendix B.", "size": 15, "color": "595959"}], pad=0.3)
notes(s, "Use for: 'Can you trust the labels?' Answer: judged supervision, not ground truth; the gold answer makes it a constrained comparison; same-model bias not measured; next step is human validation and a second judge family.")

# B7 (30) PR / ROC curves
s = content_slide("Backup: precision-recall and ROC curves on the test set")
img(s, "c7_curves", X0, 2.2, W_FULL, 4.45, "Precision-recall and ROC curves for TF-IDF and ModernBERT on the 193-question test set.")
notes(s, f"Use for: ranking quality questions. AUPRC {CMP['auprc']['tfidf']:.3f} vs {CMP['auprc']['modernbert']:.3f}; AUROC {CMP['auroc']['tfidf']:.3f} vs {CMP['auroc']['modernbert']:.3f}; neither difference significant.")

# B8 (31) statistical power
s = content_slide("Backup: why the ranking differences stay undecided")
mde = {k: 2.80 * (CMP[k]["diff_ci95"][1] - CMP[k]["diff_ci95"][0]) / 3.92 for k in ("auroc", "balanced_accuracy", "auprc", "f1_mix")}
rows_ = [["Metric", "Observed difference", "Min. detectable difference (80% power, α = 0.05)", "Test set needed"]]
for k, n_ in (("auroc", "AUROC"), ("balanced_accuracy", "Balanced accuracy"), ("auprc", "AUPRC"), ("f1_mix", "F1 (mix)")):
    rows_.append([n_, sgn(CMP[k]["diff"]), f"≈ {mde[k]:.2f}", f"≈ {(mde[k] / CMP[k]['diff'])**2:.0f}× larger"])
table(s, X0, 2.3, W_FULL, 2.6, rows_, col_w=[2.6, 2.4, 4.6, 2.42], size=17)
text(s, X0, 5.25, W_FULL, 1.4, [{"runs": f"Derived from the paired-bootstrap standard errors on {D['n_test']} test questions ({D['n_pos']} Mix-required). A non-significant result here means undecided, not equal. Required size scales with (detectable / observed)².",
                                  "size": 17, "color": DGREY, "line": 1.05}])
notes(s, "Use for: 'Is 193 test queries enough?' / 'Why not cross-validation?' (the fixed split kept threshold selection clean; CV is a fair suggestion for future work). Thesis Section 5.6.5.")

# B9 (32) cascade comparison (estimate, not in thesis)
s = content_slide("Backup: even a perfect Naive-first cascade costs more time than query-only routing")
rows_ = [["Policy", "Mean time per query", "Est. tokens per query", "Kept answers*", "Wait for Mix-needing questions"],
         ["ModernBERT router (query-only)", f"{mb['mean_time']:.1f} s", fmt_int(TOK['modernbert']), f"{mb['retention']*100:.1f}%", f"{CAS['mb_mix_latency']:.0f} s (Mix only)"],
         ["Perfect cascade: Naive first, escalate exactly the Mix-required", f"{CAS['mean_time']:.1f} s", fmt_int(round(CAS['mean_tokens'])), "100%", f"{CAS['escalated_latency']:.0f} s (Naive, then Mix)"],
         ["Always-Mix", f"{am['mean_time']:.1f} s", fmt_int(TOK['always_mix']), "100%", f"{am['mean_time']:.0f} s"]]
table(s, X0, 2.3, W_FULL, 2.4, rows_, col_w=[4.2, 1.95, 2.0, 1.6, 2.27], size=16, bold_rows=(1,))
text(s, X0, 4.95, W_FULL, 1.8, [
    {"runs": "A cascade pays for Naive on every question, and its escalated questions pay for both modes. The perfect cascade is the cheapest possible full-recall cascade; a real one needs an imperfect failure detector, so it would cost more or keep fewer answers.", "size": 17, "space_after": 8, "line": 1.05},
    {"runs": f"Estimate from the per-query times and per-mode token means on the {D['n_test']} test questions; not part of the thesis. *Kept answers under the same assumption as slide 17. Mix-needing wait: mean over the questions each policy sends to Mix.", "size": 14, "color": "595959"}])
notes(s, f"Use for: 'Why not run Naive first and escalate when it fails?' Answer: even a perfect cascade averages {CAS['mean_time']:.1f} s per query vs {mb['mean_time']:.1f} s for the query-only router, and the Mix-needing questions wait about {CAS['escalated_latency']:.0f} s because they pay for both modes. A cascade would recover answers the router misses, so it is a different point on the trade-off, and combining both (route first, escalate on admitted failure) is natural future work. Clearly say this is a back-of-the-envelope estimate, not a thesis result.")

# B10 (33) configuration
s = content_slide("Backup: experimental configuration")
rows_ = [["Component", "Setting"],
         ["Generator and judge", "gemini-3-flash-preview, temperature 0"],
         ["Embeddings", "gemini-embedding-001 (768-dim)"],
         ["LightRAG", "v1.4.10; chunks 1,200 tokens, overlap 100; Naive top_k = chunk_top_k = 5; Mix chunk_top_k = 5 (graph retrieval at LightRAG defaults)"],
         ["TF-IDF router", "uni- and bigrams, min df 2, 20k features; logistic regression, C = 1.0, balanced weights"],
         ["ModernBERT router", "ModernBERT-base, max 128 tokens, lr 2e-5, batch 16, up to 8 epochs, early stopping on validation F1"],
         ["Seeds and thresholds", f"seeds 7, 13, 42; thresholds from validation F1 (TF-IDF {D['thr']['tfidf']:.2f}, ensemble {D['thr']['modernbert']:.2f})"],
         ["Code and data", "github.com/Dantecool02/Thesis_ericsson, tag examiner-revision (labels, splits, outputs in results/)"]]
table(s, X0, 2.25, W_FULL, 4.4, rows_, col_w=[3.0, 9.02], size=16, first_col_bold=True, left_cols=(0, 1))
notes(s, "Use for configuration questions. Thesis Appendix B.")

# B11 (34) generative AI
s = content_slide("Backup: use of generative AI tools")
box(s, X0, 2.3, W_FULL, 3.4, LIGHT, [
    {"runs": "OpenAI Codex and Anthropic Claude Code were used in an advisory role, as discussion partners:", "size": 19, "space_after": 8, "line": 1.05},
    {"runs": "discussing phrasing and implementation alternatives, and debugging", "size": 19, "bullet": True, "space_after": 4},
    {"runs": "preparing the revisions after supervisor and examiner feedback, and this presentation", "size": 19, "bullet": True, "space_after": 12},
    {"runs": "The research question, design, experiments, analyses and conclusions are my own; every reported number was checked against the logs and output files.", "size": 19, "line": 1.05}], pad=0.35)
text(s, X0, 5.95, W_FULL, 0.6, [{"runs": "Thesis Section 5.7, Use of generative AI tools. The Gemini model inside the pipeline is part of the method, not of this disclosure.", "size": 15, "color": "595959"}])
notes(s, "Use for: questions about AI assistance. Keep it factual and brief.")

prs.save(OUT)
print("saved", OUT, "| slides:", len(prs.slides))
